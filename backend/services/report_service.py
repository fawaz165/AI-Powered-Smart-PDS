import os
from pathlib import Path
from datetime import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from backend.database.mongo import get_db
from backend.config import Config
from backend.utils.logger import logger

REPORTS_DIR = Config.REPORTS_DIR
REPORTS_DIR.mkdir(parents=True, exist_ok=True)


class ReportService:
    """
    Generates structured, database-backed Excel spreadsheets with support
    for district, date, commodity, and risk filtering, styled using OpenPyXL.
    """

    @staticmethod
    def _apply_excel_formatting(filepath: Path, sheet_name: str, title: str):
        """Applies professional administrative styling: dark headers, auto column width, borders."""
        try:
            wb = openpyxl.load_workbook(filepath)
            ws = wb[sheet_name]

            header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            thin_border = Border(
                left=Side(style="thin", color="CBD5E1"),
                right=Side(style="thin", color="CBD5E1"),
                top=Side(style="thin", color="CBD5E1"),
                bottom=Side(style="thin", color="CBD5E1")
            )

            # Style header row (Row 1)
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            ws.row_dimensions[1].height = 26

            # Format data rows & adjust column widths
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    cell.border = thin_border
                    val_str = str(cell.value or "")
                    if len(val_str) > max_len:
                        max_len = len(val_str)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

            wb.save(filepath)
        except Exception as e:
            logger.warning(f"Failed to apply openpyxl formatting to {filepath}: {e}")

    @staticmethod
    def _parse_date_bound(date_str: str, is_end: bool = False):
        """Parses YYYY-MM-DD string into datetime."""
        if not date_str:
            return None
        try:
            clean = str(date_str).strip()[:10]
            dt = datetime.strptime(clean, "%Y-%m-%d")
            if is_end:
                return dt.replace(hour=23, minute=59, second=59, microsecond=999999)
            return dt.replace(hour=0, minute=0, second=0, microsecond=0)
        except Exception:
            return None

    @classmethod
    def generate_inventory_report(cls, region: str = None, commodity: str = None) -> Path:
        """
        Exports full regional inventory status with stock and buffer analysis.
        Supports filtering by region and commodity.
        """
        db = get_db()
        filter_query = {}
        if region and region.upper() != "ALL":
            filter_query["region"] = region
        if commodity and commodity.upper() != "ALL":
            filter_query["commodity"] = commodity

        items = list(db.inventory.find(filter_query, {"_id": 0}))

        rows = []
        for it in items:
            cur_stock = float(it.get("current_stock", 0.0))
            buf_stock = float(it.get("buffer_stock", 0.0))
            
            if cur_stock < buf_stock:
                status = "CRITICAL SHORTAGE"
            elif cur_stock < (buf_stock * 1.5):
                status = "LOW BUFFER"
            else:
                status = "SUFFICIENT STOCK"

            cov_ratio = round(cur_stock / (buf_stock + 1e-5), 2)
            last_up = it.get("updated_at")
            if isinstance(last_up, datetime):
                up_str = last_up.strftime("%Y-%m-%d %H:%M:%S")
            else:
                up_str = str(last_up or "N/A")

            rows.append({
                "Commodity Name": it.get("commodity", ""),
                "Depot District": it.get("region", ""),
                "Current Warehouse Stock (kg)": round(cur_stock, 1),
                "Target Safety Buffer (kg)": round(buf_stock, 1),
                "Buffer Coverage Ratio": f"{cov_ratio}x",
                "Inventory Stock Status": status,
                "Last Stock Verification": up_str
            })

        df = pd.DataFrame(rows)
        if df.empty:
            df = pd.DataFrame(columns=[
                "Commodity Name", "Depot District", "Current Warehouse Stock (kg)",
                "Target Safety Buffer (kg)", "Buffer Coverage Ratio", "Inventory Stock Status",
                "Last Stock Verification"
            ])

        tag = f"_{region}" if region and region.upper() != "ALL" else ""
        filename = f"inventory_report{tag}_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = REPORTS_DIR / filename

        with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name="Warehouse Inventory", index=False)

        cls._apply_excel_formatting(filepath, "Warehouse Inventory", "Warehouse Inventory Status")
        logger.info(f"Generated inventory report: {filepath} ({len(df)} records)")
        return filepath

    @classmethod
    def generate_distribution_report(
        cls,
        region: str = None,
        commodity: str = None,
        start_date: str = None,
        end_date: str = None,
        risk_level: str = None
    ) -> Path:
        """
        Exports grain distribution transaction log with date range and district filtering.
        """
        db = get_db()
        filter_query = {}
        if region and region.upper() != "ALL":
            filter_query["region"] = region
        if commodity and commodity.upper() != "ALL":
            filter_query["commodity"] = commodity
        if risk_level and risk_level.upper() != "ALL":
            filter_query["risk_level"] = risk_level.upper()

        dt_start = cls._parse_date_bound(start_date, is_end=False)
        dt_end = cls._parse_date_bound(end_date, is_end=True)
        if dt_start or dt_end:
            time_filter = {}
            if dt_start:
                time_filter["$gte"] = dt_start
            if dt_end:
                time_filter["$lte"] = dt_end
            filter_query["timestamp"] = time_filter

        txns = list(db.transactions.find(filter_query, {"_id": 0}).sort("timestamp", -1))

        rows = []
        for t in txns:
            ts = t.get("timestamp")
            if isinstance(ts, datetime):
                utc_str = ts.strftime("%Y-%m-%d %H:%M:%S")
                local_str = ts.strftime("%d-%b-%Y %I:%M %p")
            else:
                utc_str = str(ts or "N/A")
                local_str = utc_str

            rows.append({
                "Transaction ID": t.get("transaction_id", ""),
                "Beneficiary ID": t.get("beneficiary_id", ""),
                "Ration Shop ID": t.get("ration_shop_id", ""),
                "Commodity Distributed": t.get("commodity", ""),
                "Quantity Distributed (kg)": float(t.get("quantity", 0.0)),
                "Depot District": t.get("region", ""),
                "Distribution Time (UTC)": utc_str,
                "Local Formatted Time": local_str,
                "Risk Classification": t.get("risk_level", "LOW"),
                "Isolation Forest Score": t.get("anomaly_score", 0.0),
                "Processing Status": t.get("processing_status", "COMPLETED"),
                "Officer Review Decision": t.get("review_status", "PENDING_REVIEW"),
                "Reviewing Officer": t.get("reviewed_by", "Unreviewed"),
                "Audit Rationale": t.get("reason", "")
            })

        df = pd.DataFrame(rows)
        if df.empty:
            df = pd.DataFrame(columns=[
                "Transaction ID", "Beneficiary ID", "Ration Shop ID", "Commodity Distributed",
                "Quantity Distributed (kg)", "Depot District", "Distribution Time (UTC)",
                "Local Formatted Time", "Risk Classification", "Isolation Forest Score",
                "Processing Status", "Officer Review Decision", "Reviewing Officer", "Audit Rationale"
            ])

        tag = f"_{region}" if region and region.upper() != "ALL" else ""
        filename = f"distribution_report{tag}_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = REPORTS_DIR / filename

        with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name="Transactions Ledger", index=False)

        cls._apply_excel_formatting(filepath, "Transactions Ledger", "PDS Distribution Transactions Ledger")
        logger.info(f"Generated distribution report: {filepath} ({len(df)} records)")
        return filepath

    @classmethod
    def generate_fraud_report(
        cls,
        region: str = None,
        risk_level: str = None,
        status: str = None,
        start_date: str = None,
        end_date: str = None
    ) -> Path:
        """
        Exports flagged suspicious transactions, Isolation Forest anomaly scores,
        and officer audit review decisions.
        """
        db = get_db()
        filter_query = {}
        if region and region.upper() != "ALL":
            filter_query["region"] = region
        if risk_level and risk_level.upper() != "ALL":
            filter_query["risk_level"] = risk_level.upper()
        if status and status.upper() != "ALL":
            filter_query["status"] = status.upper()

        dt_start = cls._parse_date_bound(start_date, is_end=False)
        dt_end = cls._parse_date_bound(end_date, is_end=True)
        if dt_start or dt_end:
            time_filter = {}
            if dt_start:
                time_filter["$gte"] = dt_start
            if dt_end:
                time_filter["$lte"] = dt_end
            filter_query["created_at"] = time_filter

        alerts = list(db.fraud_alerts.find(filter_query, {"_id": 0}).sort("created_at", -1))

        rows = []
        for a in alerts:
            score = float(a.get("anomaly_score", 0.0))
            score_interp = a.get("score_interpretation")
            if not score_interp:
                if score >= 0.0:
                    score_interp = "Normal distribution range"
                elif score >= -0.10:
                    score_interp = "Mild statistical deviation"
                else:
                    score_interp = "Severe outlier in tree ensemble"

            created_at = a.get("created_at")
            flagged_str = created_at.strftime("%Y-%m-%d %H:%M:%S") if isinstance(created_at, datetime) else str(created_at or "N/A")
            reviewed_at = a.get("reviewed_at")
            rev_str = reviewed_at.strftime("%Y-%m-%d %H:%M:%S") if isinstance(reviewed_at, datetime) else str(reviewed_at or "Pending")

            rows.append({
                "Alert Transaction ID": a.get("transaction_id", ""),
                "Beneficiary ID": a.get("beneficiary_id", ""),
                "Commodity": a.get("commodity", ""),
                "Withdrawn Quantity (kg)": float(a.get("quantity", 0.0)),
                "Depot District": a.get("region", ""),
                "Risk Classification": a.get("risk_level", "HIGH"),
                "Isolation Forest Decision Score": round(score, 4),
                "Outlier Score Interpretation": score_interp,
                "Card Entitlement Cap (kg)": a.get("entitled_quota", "N/A"),
                "Quantity / Quota Ratio": f"{a.get('quantity_to_quota_ratio', 'N/A')}x",
                "Explainable Detection Reasons": a.get("reason", ""),
                "Audit Review Status": a.get("status", "PENDING_REVIEW"),
                "Official Officer Decision": a.get("review_decision", "PENDING"),
                "Reviewing Officer ID": a.get("reviewed_by", "Unassigned"),
                "Officer Audit Justification": a.get("review_reason", ""),
                "Review Timestamp": rev_str,
                "Alert Logged Timestamp": flagged_str
            })

        df = pd.DataFrame(rows)
        if df.empty:
            df = pd.DataFrame(columns=[
                "Alert Transaction ID", "Beneficiary ID", "Commodity", "Withdrawn Quantity (kg)",
                "Depot District", "Risk Classification", "Isolation Forest Decision Score",
                "Outlier Score Interpretation", "Card Entitlement Cap (kg)", "Quantity / Quota Ratio",
                "Explainable Detection Reasons", "Audit Review Status", "Official Officer Decision",
                "Reviewing Officer ID", "Officer Audit Justification", "Review Timestamp", "Alert Logged Timestamp"
            ])

        tag = f"_{region}" if region and region.upper() != "ALL" else ""
        filename = f"fraud_anomaly_report{tag}_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = REPORTS_DIR / filename

        with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name="Flagged Anomalies", index=False)

        cls._apply_excel_formatting(filepath, "Flagged Anomalies", "Fraud & Statistical Anomaly Audit Log")
        logger.info(f"Generated fraud report: {filepath} ({len(df)} records)")
        return filepath

    @classmethod
    def generate_predictions_report(cls, region: str = None, commodity: str = None) -> Path:
        """
        Exports ML demand predictions and calculated procurement recommendations.
        If database collection is empty, runs baseline ML pipeline to ensure real records are present.
        """
        from backend.services.ml_service import MLService
        db = get_db()

        filter_query = {}
        if region and region.upper() != "ALL":
            filter_query["region"] = region
        if commodity and commodity.upper() != "ALL":
            filter_query["commodity"] = commodity

        preds = list(db.predictions.find(filter_query, {"_id": 0}))

        # If no audit records exist in predictions collection, generate from live inventory & ML pipeline
        if not preds:
            inv_filter = {}
            if region and region.upper() != "ALL":
                inv_filter["region"] = region
            if commodity and commodity.upper() != "ALL":
                inv_filter["commodity"] = commodity

            inv_items = list(db.inventory.find(inv_filter, {"_id": 0}))
            now = datetime.utcnow()
            for it in inv_items:
                com = it.get("commodity", "Rice")
                reg = it.get("region", "Chennai")
                stk = float(it.get("current_stock", 3800.0))
                buf = float(it.get("buffer_stock", 500.0))
                
                try:
                    p = MLService.predict_demand({
                        "commodity": com,
                        "region": reg,
                        "month": now.month,
                        "year": now.year,
                        "current_stock": stk,
                        "safety_stock": buf,
                        "previous_demand": stk * 1.1,
                        "beneficiary_count": 12500,
                        "is_custom_scenario": False
                    }, save_audit=False)
                    preds.append(p)
                except Exception as e:
                    logger.warning(f"Could not generate prediction for {com}/{reg}: {e}")

        rows = []
        for p in preds:
            bd = p.get("calculation_breakdown", {})
            created = p.get("created_at")
            created_str = created.strftime("%Y-%m-%d %H:%M:%S") if isinstance(created, datetime) else str(created or "N/A")

            rows.append({
                "Commodity": p.get("commodity", ""),
                "Depot District": p.get("region", ""),
                "Forecast Month": p.get("month", datetime.utcnow().month),
                "Forecast Year": p.get("year", datetime.utcnow().year),
                "Registered Cardholders": p.get("beneficiary_count", 12500),
                "Historical Baseline (kg)": bd.get("baseline_demand", "N/A"),
                "Seasonal Multiplier": f"{bd.get('seasonal_multiplier', 1.0):.2f}x",
                "Predicted Monthly Demand (kg)": float(p.get("predicted_demand", 0.0)),
                "Current Warehouse Stock (kg)": float(p.get("current_stock", 0.0)),
                "Target Safety Buffer (kg)": float(p.get("safety_stock", 0.0)),
                "Gross Requirement (kg)": bd.get("gross_requirement", float(p.get("predicted_demand", 0.0)) + float(p.get("safety_stock", 0.0))),
                "Net Detected Shortage (kg)": float(p.get("shortage", 0.0)),
                "Recommended Procurement (kg)": float(p.get("recommended_procurement", 0.0)),
                "Inventory Coverage Status": p.get("status", "SUFFICIENT"),
                "Procurement Formula": "max(0, (Predicted Demand + Buffer) - Stock)",
                "Production ML Model": p.get("model_used", "Linear Regression"),
                "Data Source": p.get("data_source", "DATABASE_BACKED"),
                "Generated At": created_str
            })

        df = pd.DataFrame(rows)
        if df.empty:
            df = pd.DataFrame(columns=[
                "Commodity", "Depot District", "Forecast Month", "Forecast Year",
                "Registered Cardholders", "Historical Baseline (kg)", "Seasonal Multiplier",
                "Predicted Monthly Demand (kg)", "Current Warehouse Stock (kg)", "Target Safety Buffer (kg)",
                "Gross Requirement (kg)", "Net Detected Shortage (kg)", "Recommended Procurement (kg)",
                "Inventory Coverage Status", "Procurement Formula", "Production ML Model",
                "Data Source", "Generated At"
            ])

        tag = f"_{region}" if region and region.upper() != "ALL" else ""
        filename = f"demand_predictions_report{tag}_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = REPORTS_DIR / filename

        with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name="Demand Forecast", index=False)

        cls._apply_excel_formatting(filepath, "Demand Forecast", "ML Demand Forecast & Procurement Recommendations")
        logger.info(f"Generated predictions report: {filepath} ({len(df)} records)")
        return filepath
