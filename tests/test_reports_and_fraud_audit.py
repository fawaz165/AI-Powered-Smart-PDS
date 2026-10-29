import io
import openpyxl
import pytest
from datetime import datetime, timedelta
from backend.app import create_app
from backend.database.mongo import get_db
from backend.services.ml_service import MLService


@pytest.fixture
def client():
    app = create_app()
    app.config["TESTING"] = True
    with app.test_client() as client:
        yield client


@pytest.fixture(autouse=True)
def setup_teardown_reports_and_fraud_data():
    """Sets up isolated test data for reports, fraud alerts, and entitlement checks."""
    db = get_db()

    # Clean up test records
    def cleanup():
        db.beneficiaries.delete_many({"beneficiary_id": {"$regex": "^BEN-TEST-REP"}})
        db.inventory.delete_many({"region": "ReportTestRegion"})
        db.transactions.delete_many({"region": "ReportTestRegion"})
        db.fraud_alerts.delete_many({"region": "ReportTestRegion"})

    cleanup()

    # Seed test beneficiaries
    db.beneficiaries.insert_many([
        {
            "beneficiary_id": "BEN-TEST-REP-01",
            "name": "Report PHH Beneficiary",
            "region": "ReportTestRegion",
            "family_size": 4,
            "ration_card_type": "Priority (PHH)",
            "ration_card_number": "RC-TEST-REP-001",
            "monthly_quota": {
                "Rice": 20.0,
                "Wheat": 8.0,
                "Sugar": 1.0,
                "Dal": 1.0
            }
        },
        {
            "beneficiary_id": "BEN-TEST-REP-02",
            "name": "Report AAY Beneficiary",
            "region": "ReportTestRegion",
            "family_size": 3,
            "ration_card_type": "Antyodaya (AAY)",
            "ration_card_number": "RC-TEST-REP-002",
            "monthly_quota": {
                "Rice": 25.0,
                "Wheat": 10.0,
                "Sugar": 2.0,
                "Dal": 2.0
            }
        }
    ])

    # Seed test inventory
    db.inventory.insert_many([
        {
            "commodity": "Rice",
            "commodity_code": "COM-RIC",
            "region": "ReportTestRegion",
            "current_stock": 500.0,
            "buffer_stock": 50.0,
            "last_updated": datetime.utcnow()
        },
        {
            "commodity": "Sugar",
            "commodity_code": "COM-SUG",
            "region": "ReportTestRegion",
            "current_stock": 200.0,
            "buffer_stock": 20.0,
            "last_updated": datetime.utcnow()
        }
    ])

    # Seed test transactions and fraud alert
    txn_id_1 = "TXN-REP-TEST-001"
    txn_id_2 = "TXN-REP-TEST-002"

    db.transactions.insert_many([
        {
            "transaction_id": txn_id_1,
            "beneficiary_id": "BEN-TEST-REP-01",
            "commodity": "Rice",
            "quantity": 20.0,
            "region": "ReportTestRegion",
            "ration_shop_id": "FPS-REP-001",
            "timestamp": datetime.utcnow() - timedelta(days=2),
            "processing_status": "COMPLETED",
            "previous_stock": 520.0,
            "remaining_stock": 500.0,
            "audit_trail": [
                {"action": "TRANSACTION_CREATED", "timestamp": (datetime.utcnow() - timedelta(days=2)).isoformat()}
            ]
        },
        {
            "transaction_id": txn_id_2,
            "beneficiary_id": "BEN-TEST-REP-01",
            "commodity": "Sugar",
            "quantity": 10.0,  # 10x entitlement quota!
            "region": "ReportTestRegion",
            "ration_shop_id": "FPS-REP-001",
            "timestamp": datetime.utcnow() - timedelta(days=1),
            "processing_status": "COMPLETED",
            "previous_stock": 210.0,
            "remaining_stock": 200.0,
            "anomaly_score": -0.15,
            "risk_level": "HIGH",
            "audit_trail": [
                {"action": "TRANSACTION_CREATED", "timestamp": (datetime.utcnow() - timedelta(days=1)).isoformat()}
            ]
        }
    ])

    db.fraud_alerts.insert_one({
        "alert_id": "ALT-REP-TEST-001",
        "transaction_id": txn_id_2,
        "beneficiary_id": "BEN-TEST-REP-01",
        "region": "ReportTestRegion",
        "commodity": "Sugar",
        "quantity": 10.0,
        "entitled_quota": 1.0,
        "quantity_to_quota_ratio": 10.0,
        "anomaly_score": -0.15,
        "score_interpretation": "Isolation Forest identified this pattern as severe statistical outlier",
        "score_badge": "SEVERE_ANOMALY",
        "risk_level": "HIGH",
        "reason": "Distribution quantity 10.00 kg exceeds card quota 1.00 kg (Ratio: 10.0x)",
        "status": "PENDING_REVIEW",
        "timestamp": datetime.utcnow() - timedelta(days=1),
        "review_history": []
    })

    yield

    cleanup()


def test_excel_inventory_report_export(client):
    """Verifies that GET /api/reports/inventory returns a valid, populated XLSX file."""
    res = client.get("/api/reports/inventory?region=ReportTestRegion")
    assert res.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in res.content_type

    # Validate openpyxl loads the XLSX file properly
    wb = openpyxl.load_workbook(io.BytesIO(res.data))
    assert "Warehouse Inventory" in wb.sheetnames
    ws = wb["Warehouse Inventory"]

    # Verify headers
    headers = [cell.value for cell in ws[1]]
    assert "Depot District" in headers
    assert "Commodity Name" in headers
    assert "Current Warehouse Stock (kg)" in headers

    # Verify rows contain our test records
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) >= 3  # Header + 2 items
    regions = [row[1] for row in rows[1:]]  # Depot District is column 2
    assert all(r == "ReportTestRegion" for r in regions)


def test_excel_distribution_report_export_with_date_filters(client):
    """Verifies that GET /api/reports/distribution returns valid XLSX with date and district filtering."""
    start = (datetime.utcnow() - timedelta(days=5)).strftime("%Y-%m-%d")
    end = (datetime.utcnow() + timedelta(days=1)).strftime("%Y-%m-%d")

    res = client.get(f"/api/reports/distribution?region=ReportTestRegion&commodity=Rice&start_date={start}&end_date={end}")
    assert res.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in res.content_type

    wb = openpyxl.load_workbook(io.BytesIO(res.data))
    assert "Transactions Ledger" in wb.sheetnames
    ws = wb["Transactions Ledger"]

    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) >= 2  # Header + Rice txn
    # First row is header
    commodities = [row[3] for row in rows[1:]]  # Commodity Distributed is index 3
    assert all(c == "Rice" for c in commodities)


def test_excel_fraud_report_export(client):
    """Verifies that GET /api/reports/fraud exports real database fraud alerts with score interpretations."""
    res = client.get("/api/reports/fraud?region=ReportTestRegion&risk_level=HIGH")
    assert res.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in res.content_type

    wb = openpyxl.load_workbook(io.BytesIO(res.data))
    assert "Flagged Anomalies" in wb.sheetnames
    ws = wb["Flagged Anomalies"]

    headers = [cell.value for cell in ws[1]]
    assert "Alert Transaction ID" in headers
    assert "Isolation Forest Decision Score" in headers
    assert "Outlier Score Interpretation" in headers
    assert "Audit Review Status" in headers

    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) >= 2  # Header + at least ALT-REP-TEST-001
    alert_ids = [row[0] for row in rows[1:]]
    assert "TXN-REP-TEST-002" in alert_ids


def test_excel_predictions_report_export(client):
    """Verifies that GET /api/reports/predictions exports valid XLSX backed by inventory & ML pipeline."""
    res = client.get("/api/reports/predictions?region=ReportTestRegion")
    assert res.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in res.content_type

    wb = openpyxl.load_workbook(io.BytesIO(res.data))
    assert "Demand Forecast" in wb.sheetnames
    ws = wb["Demand Forecast"]

    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) >= 2  # Header + at least one prediction row
    headers = [cell.value for cell in ws[1]]
    assert "Commodity" in headers
    assert "Predicted Monthly Demand (kg)" in headers


def test_commodity_specific_entitlement_validation():
    """Verifies MLService validates commodity-specific quotas (e.g. Sugar 1kg vs Rice 20kg)."""
    # 1. Sugar draw of 10.0 kg for a PHH family (quota is 1.0 kg) -> 10x ratio -> MUST be flagged SUSPICIOUS
    sugar_eval = MLService.detect_anomaly(
        beneficiary_id="BEN-TEST-REP-01",
        commodity="Sugar",
        quantity=10.0
    )
    assert sugar_eval["status"] == "SUSPICIOUS"
    assert sugar_eval["risk_level"] in ["HIGH", "MEDIUM"]
    assert sugar_eval["requires_admin_review"] is True
    assert "1.0 kg" in sugar_eval["reason"]
    assert "score_interpretation" in sugar_eval
    assert "score_badge" in sugar_eval
    assert "audit_disclaimer" in sugar_eval

    # 2. Rice draw of 20.0 kg for the same PHH family (quota is 20.0 kg) -> 1x ratio -> MUST be NORMAL
    rice_eval = MLService.detect_anomaly(
        beneficiary_id="BEN-TEST-REP-01",
        commodity="Rice",
        quantity=20.0
    )
    assert rice_eval["status"] == "NORMAL"
    assert rice_eval["risk_level"] == "LOW"
    assert rice_eval["requires_admin_review"] is False
    assert rice_eval["score_badge"] in ["NORMAL_INLIER", "BORDERLINE_NORMAL"]


def test_fraud_alert_review_audit_workflow(client):
    """Verifies official review decision workflow with officer ID, timestamp, and audit reason."""
    db = get_db()
    txn_id = "TXN-REP-TEST-002"

    # 1. Validation: Missing officer ID
    res = client.post(f"/api/transactions/fraud-alerts/{txn_id}/review", json={
        "decision": "APPROVED_GENUINE",
        "reason": "Special family festival allocation authorized by civil supplies director."
    })
    assert res.status_code == 400
    assert "Officer ID is required" in res.get_json()["message"]

    # 2. Validation: Invalid decision enum
    res = client.post(f"/api/transactions/fraud-alerts/{txn_id}/review", json={
        "officer_id": "OFFICER-007",
        "decision": "INVALID_DECISION_TYPE",
        "reason": "Random reason"
    })
    assert res.status_code == 400
    assert "Invalid review decision" in res.get_json()["message"]

    # 3. Validation: Missing / empty reason
    res = client.post(f"/api/transactions/fraud-alerts/{txn_id}/review", json={
        "officer_id": "OFFICER-007",
        "decision": "FLAGGED_FRAUD",
        "reason": "  "
    })
    assert res.status_code == 400
    assert "Audit reason / rationale is required" in res.get_json()["message"]

    # 4. Valid Review Submission: APPROVED_GENUINE
    valid_payload = {
        "officer_id": "INSPECTOR-TAMIL-042",
        "decision": "APPROVED_GENUINE",
        "reason": "Verified genuine requirement during Pongal festival special entitlement quota."
    }
    res = client.post(f"/api/transactions/fraud-alerts/{txn_id}/review", json=valid_payload)
    assert res.status_code == 200
    res_data = res.get_json()
    assert res_data["status"] == "success"
    assert "Audit review for" in res_data["message"]

    # Verify fraud alert was updated with review history
    alert = db.fraud_alerts.find_one({"transaction_id": txn_id})
    assert alert is not None
    assert alert["status"] == "REVIEWED"
    assert alert["review_decision"] == "APPROVED_GENUINE"
    assert alert["reviewed_by"] == "INSPECTOR-TAMIL-042"
    assert "Pongal festival" in alert["review_reason"]
    assert len(alert["review_history"]) == 1
    assert alert["review_history"][0]["decision"] == "APPROVED_GENUINE"

    # Verify original transaction was preserved and NOT mutated
    txn = db.transactions.find_one({"transaction_id": txn_id})
    assert txn is not None
    assert txn["quantity"] == 10.0  # Original transaction quantity untouched
    assert txn["remaining_stock"] == 200.0  # Inventory balance untouched
    assert txn["anomaly_score"] == -0.15  # ML anomaly score preserved

    # Verify audit trail appended an administrative review event
    audit_events = [ev for ev in txn.get("audit_trail", []) if ev.get("action") == "ADMINISTRATIVE_FRAUD_REVIEW"]
    assert len(audit_events) == 1
    assert audit_events[0]["officer_id"] == "INSPECTOR-TAMIL-042"
    assert audit_events[0]["decision"] == "APPROVED_GENUINE"


def test_backward_compatible_resolve_endpoint(client):
    """Verifies that legacy /resolve endpoint forwards to the review workflow."""
    txn_id = "TXN-REP-TEST-002"
    payload = {
        "officer_id": "OFFICER-LEGACY-01",
        "decision": "DISMISSED_FALSE_POSITIVE",
        "reason": "Data entry typo corrected by regional supervisor.",
        "resolution_notes": "Forwarded from legacy interface"
    }

    res = client.post(f"/api/transactions/fraud-alerts/{txn_id}/resolve", json=payload)
    assert res.status_code == 200
    data = res.get_json()
    assert data["status"] == "success"
